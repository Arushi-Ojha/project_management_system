import os
from docx import Document
from docx.shared import Inches
from datetime import date

def create_report():
    doc = Document()
    
    # Title
    doc.add_heading('Project Management System - Comprehensive Report', 0)
    
    # Date
    today = date.today().strftime("%B %d, %Y")
    doc.add_paragraph(f"Date: {today}")
    
    # Links
    doc.add_heading('1. Important Links', level=1)
    doc.add_paragraph("Deployed Frontend URL: https://proman-aura-frontend-362935833196.asia-south1.run.app")
    doc.add_paragraph("Deployed Backend API URL: https://proman-aura-backend-362935833196.asia-south1.run.app")
    doc.add_paragraph("GitHub Repository: https://github.com/Arushi-Ojha/project_management_system")

    # Features
    doc.add_heading('2. Features', level=1)
    features = [
        "Authentication & Security: Secure JWT-based authentication with 6-digit OTP verification. Includes honeypot spam protection.",
        "Role-Based Access Control (RBAC): Strict data isolation for Organizations (Owners), Admins, and Members.",
        "Mass User Onboarding: Bulk CSV upload to onboard users, automatic temporary password generation, and email notifications.",
        "Dynamic Role-based Dashboards: Tailored views depending on user roles, including Analytics for owners and Kanban boards for members.",
        "System Audit Trail: Comprehensive logging of system events with RBAC visibility.",
        "Kanban Boards: Interactive drag-and-drop boards with assignee tracking, deadline alerts, and automatic email notifications.",
        "Mobile-Friendly Aesthetics: Fully responsive design, utilizing flexible grids and removing clutter on mobile devices."
    ]
    for feature in features:
        doc.add_paragraph(f"- {feature}", style='List Bullet')
        
    # API Endpoints
    doc.add_heading('3. API Endpoints', level=1)
    endpoints = [
        "Auth: POST /api/v1/iam/auth/signup, POST /api/v1/iam/auth/verify-otp, POST /api/v1/iam/auth/login",
        "Users: POST /api/v1/iam/users, GET /api/v1/iam/organizations/{org_id}/users, GET /api/v1/iam/users/lookup",
        "Mass Actions: GET /api/v1/iam/organizations/{org_id}/users/export, POST /api/v1/iam/organizations/{org_id}/users/bulk-upload",
        "Teams: POST /api/v1/iam/teams, GET /api/v1/iam/organizations/{org_id}/teams, PATCH /api/v1/iam/teams/{id}, DELETE /api/v1/iam/teams/{id}",
        "Issues: POST /api/v1/issues/, GET /api/v1/issues/, GET /api/v1/issues/{id}, PATCH /api/v1/issues/{id}",
        "Extensions: POST /api/v1/extensions/webhooks, GET /api/v1/extensions/audit-logs"
    ]
    for ep in endpoints:
        doc.add_paragraph(f"- {ep}", style='List Bullet')
        
    # Pipeline and Workflow
    doc.add_heading('4. Deployment Pipeline & Workflow', level=1)
    doc.add_paragraph(
        "The application utilizes a containerized microservice architecture deployed on Google Cloud Run:\n"
        "1. Backend: A FastAPI Python backend connected to MongoDB Atlas. It is containerized using Docker, exposing port 8080. It utilizes PyJWT for auth and Motor for async database operations.\n"
        "2. Frontend: A React Vite application. The Dockerfile utilizes a multi-stage build, compiling the React app and serving the static files via an Nginx web server configured for SPA routing and aggressive caching.\n"
        "3. Workflow: Code pushed to the repository represents the production state. Cloud Run automatically scales the stateless containers from zero to handle incoming traffic securely over HTTPS, leveraging strict CORS rules."
    )
    
    # Screenshots
    doc.add_heading('5. Application Screenshots', level=1)
    ss_dir = 'ss'
    if os.path.exists(ss_dir):
        for img_name in sorted(os.listdir(ss_dir)):
            if img_name.lower().endswith(('.png', '.jpg', '.jpeg')):
                img_path = os.path.join(ss_dir, img_name)
                doc.add_paragraph(f"Screenshot: {img_name}")
                try:
                    doc.add_picture(img_path, width=Inches(6.0))
                except Exception as e:
                    doc.add_paragraph(f"(Error loading image: {str(e)})")
    else:
        doc.add_paragraph("No screenshots found.")
        
    doc.save('Project_Report.docx')
    print("Report generated successfully: Project_Report.docx")

if __name__ == '__main__':
    create_report()
